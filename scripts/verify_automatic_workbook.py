from __future__ import annotations

import json
import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


EXPECTED_SHEETS = [
    "Portfolios",
    "Sponsored Products Campaigns",
    "Sponsored Display Campaigns",
    "Sponsored Brands Campaigns",
    "SB Multi Ad Group Campaigns",
    "RAS Campaigns",
    "Config",
]
BASE_ENTITIES = ["Campaign", "Ad Group", "Product Ad", "Product Targeting"]
EXPECTED_TARGETS = {"close-match", "loose-match", "substitutes", "complements"}


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("Usage: verify_automatic_workbook.py <generated.xlsx>")

    path = Path(sys.argv[1]).resolve()
    workbook = load_workbook(path, data_only=False, read_only=False)
    assert workbook.sheetnames == EXPECTED_SHEETS, workbook.sheetnames
    assert workbook["Config"].sheet_state == "veryHidden"

    sheet = workbook["Sponsored Products Campaigns"]
    assert sheet.max_column == 32, sheet.max_column
    row_numbers = range(2, sheet.max_row + 1)
    rows_by_entity = {}
    for row in row_numbers:
        entity = sheet.cell(row=row, column=2).value
        rows_by_entity.setdefault(entity, []).append(row)

    campaign_rows = rows_by_entity.get("Campaign", [])
    campaign_count = len(campaign_rows)
    assert campaign_count > 0
    for entity in BASE_ENTITIES[:2]:
        assert len(rows_by_entity.get(entity, [])) == campaign_count, (entity, rows_by_entity)

    campaign_ids = [sheet.cell(row=row, column=4).value for row in campaign_rows]
    campaign_id_set = set(campaign_ids)
    targeting_types = [sheet.cell(row=row, column=14).value for row in campaign_rows]
    budgets = [sheet.cell(row=row, column=16).value for row in campaign_rows]
    skus = [sheet.cell(row=row, column=17).value for row in rows_by_entity["Product Ad"]]
    default_bids = {
        sheet.cell(row=row, column=4).value: sheet.cell(row=row, column=18).value
        for row in rows_by_entity["Ad Group"]
    }
    target_bids = [sheet.cell(row=row, column=19).value for row in rows_by_entity["Product Targeting"]]
    expressions = [
        sheet.cell(row=row, column=27).value
        for row in rows_by_entity["Product Targeting"]
    ]
    negative_keyword_rows = rows_by_entity.get("Negative Keyword", [])
    negative_product_rows = rows_by_entity.get("Negative Product Targeting", [])

    assert len(campaign_id_set) == campaign_count
    assert set(targeting_types) == {"AUTO"}
    assert set(expressions).issubset(EXPECTED_TARGETS) and expressions
    assert all(isinstance(value, (int, float)) and value > 0 for value in budgets)
    assert len(set(budgets)) == 1, budgets
    assert all(isinstance(value, (int, float)) and value > 0 for value in default_bids.values())
    assert all(isinstance(value, (int, float)) and value > 0 for value in target_bids)
    assert all(skus)
    products_by_campaign = defaultdict(set)
    targets_by_campaign = defaultdict(set)
    for row in rows_by_entity["Product Ad"]:
        products_by_campaign[sheet.cell(row=row, column=4).value].add(sheet.cell(row=row, column=17).value)
    for row in rows_by_entity["Product Targeting"]:
        campaign_id = sheet.cell(row=row, column=4).value
        targets_by_campaign[campaign_id].add(sheet.cell(row=row, column=27).value)
        assert sheet.cell(row=row, column=19).value == default_bids[campaign_id]
    assert set(products_by_campaign) == campaign_id_set
    assert set(targets_by_campaign) == campaign_id_set
    assert len({frozenset(values) for values in products_by_campaign.values()}) == 1
    assert len({frozenset(values) for values in targets_by_campaign.values()}) == 1
    for entity in BASE_ENTITIES[1:]:
        assert {
            sheet.cell(row=row, column=4).value
            for row in rows_by_entity[entity]
        } == campaign_id_set
    assert len(negative_keyword_rows) % campaign_count == 0
    assert len(negative_product_rows) % campaign_count == 0
    assert all(
        sheet.cell(row=row, column=23).value in {"negativeExact", "negativePhrase"}
        for row in negative_keyword_rows
    )
    assert all(
        re.fullmatch(r'asin="[A-Z0-9]{10}"', str(sheet.cell(row=row, column=27).value))
        for row in negative_product_rows
    )
    assert all(
        sheet.cell(row=row, column=4).value in campaign_id_set
        and sheet.cell(row=row, column=5).value
        for row in [*negative_keyword_rows, *negative_product_rows]
    )

    entities = [sheet.cell(row=row, column=2).value for row in row_numbers]

    print(json.dumps({
        "ok": True,
        "path": str(path),
        "sheets": workbook.sheetnames,
        "config_state": workbook["Config"].sheet_state,
        "sp_rows": sheet.max_row - 1,
        "sp_columns": sheet.max_column,
        "campaign_count": campaign_count,
        "entity_counts": {
            entity: entities.count(entity)
            for entity in [*BASE_ENTITIES, "Negative Keyword", "Negative Product Targeting"]
        },
        "targeting_type": "AUTO",
        "targeting_expressions": sorted(set(expressions)),
        "skus": sorted(set(skus)),
        "daily_budget_total": round(sum(budgets), 2),
        "daily_budget_per_campaign": budgets[0],
        "planned_bids": [default_bids[campaign_id] for campaign_id in campaign_ids],
        "negative_keywords_per_campaign": len(negative_keyword_rows) // campaign_count,
        "negative_products_per_campaign": len(negative_product_rows) // campaign_count,
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
