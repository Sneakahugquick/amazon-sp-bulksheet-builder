from __future__ import annotations

import json
import sys
import time
from pathlib import Path

from openpyxl import load_workbook


EXPECTED_SHEETS = [
    "Portfolios",
    "Sponsored Products Campaigns",
    "Sponsored Display Campaigns",
    "Sponsored Brands Campaigns",
    "SB Multi Ad Group Campaigns",
    "RAS Campaigns",
    "Config",
]


def main() -> None:
    time.sleep(1)
    if len(sys.argv) != 2:
        raise SystemExit("Usage: verify_generated_workbook.py <generated.xlsx>")

    path = Path(sys.argv[1]).resolve()
    workbook = load_workbook(path, data_only=False, read_only=False)
    assert workbook.sheetnames == EXPECTED_SHEETS, workbook.sheetnames
    assert workbook["Config"].sheet_state == "veryHidden"

    sheet = workbook["Sponsored Products Campaigns"]
    assert sheet.max_column == 32, sheet.max_column
    assert sheet.max_row == 13, sheet.max_row

    entities = [sheet.cell(row=row, column=2).value for row in range(2, 14)]
    assert entities == [
        "Campaign", "Ad Group", "Product Ad", "Keyword",
        "Campaign", "Ad Group", "Product Ad", "Keyword",
        "Campaign", "Ad Group", "Product Ad", "Keyword",
    ], entities

    expected_names = [
        "abstract wall art-exact",
        "neutral wall decor-phrase",
        "large canvas art-broad",
    ]
    campaign_ids = [sheet.cell(row=row, column=4).value for row in (2, 6, 10)]
    campaign_names = [sheet.cell(row=row, column=10).value for row in (2, 6, 10)]
    ad_group_ids = [sheet.cell(row=row, column=5).value for row in (3, 7, 11)]
    ad_group_names = [sheet.cell(row=row, column=11).value for row in (3, 7, 11)]
    assert campaign_ids == expected_names, campaign_ids
    assert campaign_names == expected_names, campaign_names
    assert ad_group_ids == expected_names, ad_group_ids
    assert ad_group_names == expected_names, ad_group_names
    assert [sheet.cell(row=row, column=17).value for row in (4, 8, 12)] == [
        "CANVAS-SKU-001",
        "CANVAS-SKU-001",
        "CANVAS-SKU-001",
    ]

    keywords = [sheet.cell(row=row, column=20).value for row in (5, 9, 13)]
    match_types = [sheet.cell(row=row, column=23).value for row in (5, 9, 13)]
    bids = [sheet.cell(row=row, column=19).value for row in (5, 9, 13)]
    assert keywords == ["abstract wall art", "neutral wall decor", "large canvas art"]
    assert match_types == ["exact", "phrase", "broad"]
    assert bids == [0.68, 0.54, 0.42]
    assert all(isinstance(value, (int, float)) for value in bids)
    assert isinstance(sheet["P2"].value, (int, float))
    assert all(isinstance(sheet.cell(row=row, column=18).value, (int, float)) for row in (3, 7, 11))
    assert [sheet.cell(row=row, column=12).value for row in (2, 6, 10)] == [
        "20260901",
        "20260901",
        "20260901",
    ]

    print(
        json.dumps(
            {
                "ok": True,
                "path": str(path),
                "sheets": workbook.sheetnames,
                "config_state": workbook["Config"].sheet_state,
                "sp_rows": sheet.max_row - 1,
                "sp_columns": sheet.max_column,
                "entities": entities,
                "campaign_ids": campaign_ids,
                "campaign_names": campaign_names,
                "ad_group_ids": ad_group_ids,
                "ad_group_names": ad_group_names,
                "sku": "CANVAS-SKU-001",
                "keywords": keywords,
                "match_types": match_types,
                "bids": bids,
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
