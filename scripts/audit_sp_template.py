from __future__ import annotations

import json
import sys
import time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def main() -> None:
    time.sleep(1)
    if len(sys.argv) not in {2, 3}:
        raise SystemExit(
            "Usage: audit_sp_template.py <workbook.xlsx> [--conditional-only]"
        )

    workbook_path = Path(sys.argv[1]).resolve()
    workbook = load_workbook(workbook_path, data_only=False, read_only=False)
    worksheet = workbook["Sponsored Products Campaigns"]
    config = workbook["Config"]

    headers = []
    for column_index in range(1, worksheet.max_column + 1):
        cell = worksheet.cell(row=1, column=column_index)
        headers.append(
            {
                "column": get_column_letter(column_index),
                "header": cell.value,
                "style_id": cell.style_id,
                "width": worksheet.column_dimensions[get_column_letter(column_index)].width,
            }
        )

    validations = []
    for item in worksheet.data_validations.dataValidation:
        validations.append(
            {
                "type": item.type,
                "sqref": str(item.sqref),
                "formula1": item.formula1,
                "allow_blank": item.allow_blank,
                "error_title": item.errorTitle,
                "error": item.error,
                "prompt_title": item.promptTitle,
                "prompt": item.prompt,
            }
        )

    conditional_formats = []
    for item in worksheet.conditional_formatting:
        conditional_formats.append(
            {
                "sqref": str(item.sqref),
                "rules": [
                    {
                        "type": rule.type,
                        "operator": rule.operator,
                        "formula": rule.formula,
                        "dxf_id": rule.dxfId,
                    }
                    for rule in worksheet.conditional_formatting[item]
                ],
            }
        )

    config_rows = []
    for row_index in range(1, config.max_row + 1):
        values = [
            config.cell(row=row_index, column=column_index).value
            for column_index in range(1, config.max_column + 1)
        ]
        while values and values[-1] is None:
            values.pop()
        config_rows.append({"row": row_index, "values": values})

    report: dict[str, Any] = {
        "path": str(workbook_path),
        "sheet_inventory": [
            {
                "title": sheet.title,
                "state": sheet.sheet_state,
                "dimensions": sheet.calculate_dimension(),
            }
            for sheet in workbook.worksheets
        ],
        "sp_sheet": {
            "title": worksheet.title,
            "freeze_panes": str(worksheet.freeze_panes),
            "headers": headers,
            "data_validations": validations,
            "conditional_formats": conditional_formats,
        },
        "config_rows": config_rows,
    }

    if len(sys.argv) == 3 and sys.argv[2] == "--conditional-only":
        differential_styles = []
        for index, style in enumerate(workbook._differential_styles.styles):
            fill = style.fill
            differential_styles.append(
                {
                    "dxf_id": index,
                    "fill_type": fill.fill_type if fill else None,
                    "fg_color": {
                        "type": str(fill.fgColor.type) if fill else None,
                        "rgb": str(fill.fgColor.rgb) if fill else None,
                        "indexed": str(fill.fgColor.indexed) if fill else None,
                    },
                }
            )
        print(
            json.dumps(
                {
                    "conditional_formats": conditional_formats,
                    "differential_styles": differential_styles,
                },
                indent=2,
                ensure_ascii=False,
            )
        )
        return

    print(json.dumps(report, indent=2, ensure_ascii=False, default=str))


if __name__ == "__main__":
    main()
