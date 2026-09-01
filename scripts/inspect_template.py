from __future__ import annotations

import json
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def serialize(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("Usage: inspect_template.py <workbook.xlsx>")

    workbook_path = Path(sys.argv[1]).resolve()
    workbook = load_workbook(workbook_path, data_only=False, read_only=False)

    report: dict[str, Any] = {
        "path": str(workbook_path),
        "sheet_names": workbook.sheetnames,
        "defined_names": [
            {
                "name": item.name,
                "attr_text": item.attr_text,
                "hidden": item.hidden,
            }
            for item in workbook.defined_names.values()
        ],
        "sheets": [],
    }

    for worksheet in workbook.worksheets:
        populated_cells: list[dict[str, Any]] = []
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                populated_cells.append(
                    {
                        "coordinate": cell.coordinate,
                        "value": serialize(cell.value),
                        "data_type": cell.data_type,
                        "number_format": cell.number_format,
                    }
                )

        validations = []
        if worksheet.data_validations is not None:
            for validation in worksheet.data_validations.dataValidation:
                validations.append(
                    {
                        "type": validation.type,
                        "sqref": str(validation.sqref),
                        "formula1": validation.formula1,
                        "formula2": validation.formula2,
                        "allow_blank": validation.allow_blank,
                    }
                )

        tables = [
            {
                "name": table.name,
                "display_name": table.displayName,
                "ref": table.ref,
            }
            for table in worksheet.tables.values()
        ]

        hidden_columns = [
            key
            for key, dimension in worksheet.column_dimensions.items()
            if dimension.hidden
        ]
        hidden_rows = [
            index
            for index, dimension in worksheet.row_dimensions.items()
            if dimension.hidden
        ]

        report["sheets"].append(
            {
                "title": worksheet.title,
                "state": worksheet.sheet_state,
                "dimensions": worksheet.calculate_dimension(),
                "max_row": worksheet.max_row,
                "max_column": worksheet.max_column,
                "freeze_panes": str(worksheet.freeze_panes or ""),
                "auto_filter": worksheet.auto_filter.ref,
                "merged_ranges": [str(item) for item in worksheet.merged_cells.ranges],
                "hidden_columns": hidden_columns,
                "hidden_rows": hidden_rows,
                "tables": tables,
                "data_validations": validations,
                "populated_cells": populated_cells,
            }
        )

    print(json.dumps(report, indent=2, ensure_ascii=False, default=str))


if __name__ == "__main__":
    main()
